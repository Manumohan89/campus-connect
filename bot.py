import telebot
from telebot import types
import requests
import os
import psycopg2
from psycopg2 import pool
import openpyxl
from dotenv import load_dotenv
import bcrypt
import tempfile
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib.units import inch
from PIL import Image, ImageDraw
import logging
from apscheduler.schedulers.background import BackgroundScheduler
from apscheduler.triggers.cron import CronTrigger
from uuid import uuid4
import time
import aspose.pdf as ap
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

load_dotenv()

# Initialize PostgreSQL connection pool
DATABASE_URL = os.getenv('DATABASE_URL')
db_pool = pool.SimpleConnectionPool(1, 10, DATABASE_URL)

def get_db_connection():
    try:
        conn = db_pool.getconn()
        return conn
    except Exception as e:
        print(f"Error connecting to the database: {e}")
        return None

def close_db_connection(conn):
    try:
        if conn:
            db_pool.putconn(conn)
    except Exception as e:
        print(f"Error closing the database connection: {e}")

# Create tables function
def create_tables():
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    user_id SERIAL PRIMARY KEY,
                    full_name TEXT,
                    username TEXT UNIQUE,
                    password BYTEA,
                    semester TEXT,
                    college TEXT,
                    mobile TEXT CHECK (length(mobile) = 10),
                    branch TEXT,
                    year_scheme TEXT,
                    sgpa REAL,
                    cgpa REAL,
                    chat_id TEXT
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS marks (
                    mark_id SERIAL PRIMARY KEY,
                    user_id INTEGER,
                    subject_code TEXT,
                    subject_name TEXT,
                    internal_marks INTEGER,
                    external_marks INTEGER,
                    total INTEGER,
                    sgpa REAL,
                    credits INTEGER,
                    updated_on TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(user_id) REFERENCES users(user_id)
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS marks_cards (
                    card_id SERIAL PRIMARY KEY,
                    user_id INTEGER,
                    file_id TEXT,
                    uploaded_on TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(user_id) REFERENCES users(user_id)
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS reminders (
                    reminder_id SERIAL PRIMARY KEY,
                    user_id INTEGER,
                    time_str TEXT,
                    message TEXT,
                    job_id TEXT,
                    FOREIGN KEY(user_id) REFERENCES users(user_id)
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS job_opportunities (
                    job_id SERIAL PRIMARY KEY,
                    title TEXT,
                    company TEXT,
                    link TEXT,
                    description TEXT
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS feedback (
                    feedback_id SERIAL PRIMARY KEY,
                    user_id INTEGER,
                    feedback_text TEXT,
                    submitted_on TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(user_id) REFERENCES users(user_id)
                )
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS shared_documents (
                    doc_id SERIAL PRIMARY KEY,
                    user_id INTEGER,
                    file_id TEXT,
                    file_name TEXT,
                    mime_type TEXT,
                    FOREIGN KEY(user_id) REFERENCES users(user_id)
                )
            """)
            conn.commit()
            print('Tables created successfully.')
        except Exception as e:
            print(f"Error creating tables: {e}")
        finally:
            cur.close()
            close_db_connection(conn)
    else:
        print('Failed to connect to the database.')

# Call the function to create tables when bot starts
create_tables()

# Initialize bot
BOT_TOKEN = os.getenv('BOT_TOKEN')
bot = telebot.TeleBot(BOT_TOKEN)

# States for user registration and login
states = {
    'USERNAME': 0,
    'PASSWORD': 1,
    'FULL_NAME': 2,
    'LOGIN_USERNAME': 3,
    'LOGIN_PASSWORD': 4,
    'MARKSCARD_PDF': 5,
    'SEMESTER': 6,
    'COLLEGE': 7,
    'MOBILE': 8,
    'BRANCH': 9,
    'YEAR_SCHEME': 10,
    'RESET_PASSWORD': 11,
    'UPDATE_PROFILE': 12,
    'UPDATE_PROFILE_FIELD': 13,  # Ensure this state is correctly defined
    'REMINDER_TIME': 14,
    'REMINDER_MESSAGE': 15,
    'FEEDBACK': 16,
    'SHARE_DOCUMENT': 17,
}

user_sessions = {}

def init_session(chat_id):
    if chat_id not in user_sessions:
        user_sessions[chat_id] = {'state': None, 'username': None, 'userId': None}

def hash_password(password):
    salt = bcrypt.gensalt()
    hashed = bcrypt.hashpw(password.encode('utf-8'), salt)
    return hashed

def check_password(stored_password, provided_password):
    return bcrypt.checkpw(provided_password.encode('utf-8'), stored_password)

@bot.message_handler(commands=['start'])
def handle_start(message):
    init_session(message.chat.id)
    
    # Sending image
    with open('start.jpg', 'rb') as image:
        bot.send_photo(message.chat.id, image, caption="Welcome to the Student Bot!")
    
    description = """
    This bot helps you manage your student information. You can:
    - Register and login
    - Upload your marks card
    - Check your SGPA and CGPA
    - View and update your profile
    - Set and manage reminders
    - Generate and download your SGPA/CGPA report
    - Get internship/job opportunities
    - Share resources
    - Give feedback
    """
    bot.send_message(message.chat.id, description)
    
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True)
    markup.add(types.KeyboardButton('Menu'))
    bot.send_message(message.chat.id, 'Use the button below to navigate the menu:', reply_markup=markup)

@bot.message_handler(func=lambda message: message.text == 'Menu')
def handle_menu(message):
    markup = types.InlineKeyboardMarkup(row_width=2)
    markup.add(types.InlineKeyboardButton("Register", callback_data='register'),
               types.InlineKeyboardButton("Login", callback_data='login'),
               types.InlineKeyboardButton("Upload Marks Card PDF", callback_data='upload_markscard_pdf'),
               types.InlineKeyboardButton("SGPA", callback_data='sgpa'),
               types.InlineKeyboardButton("CGPA", callback_data='cgpa'),
               types.InlineKeyboardButton("Profile", callback_data='profile'),
               types.InlineKeyboardButton("Update Profile", callback_data='update_profile'),
               types.InlineKeyboardButton("Generate Report", callback_data='generate_report'),
               types.InlineKeyboardButton("Set Reminder", callback_data='set_reminder'),
               types.InlineKeyboardButton("Share Document", callback_data='share_document'),
               types.InlineKeyboardButton("List Resources", callback_data='list_resources'),
               types.InlineKeyboardButton("Job Opportunities", callback_data='job_opportunities'),
               types.InlineKeyboardButton("Feedback", callback_data='feedback'),
               types.InlineKeyboardButton("Logout", callback_data='logout'))
    bot.send_message(message.chat.id, 'Use the menu below to navigate:', reply_markup=markup)

@bot.callback_query_handler(func=lambda call: True)
def handle_query(call):
    chat_id = call.message.chat.id
    user_id = user_sessions[chat_id]['userId']

    if call.data == 'register':
        if user_id:
            bot.send_message(chat_id, 'Please logout first using /logout before registering a new account.')
        else:
            handle_register(call.message)
    elif call.data == 'login':
        if user_id:
            bot.send_message(chat_id, 'Please logout first using /logout before logging in.')
        else:
            handle_login(call.message)
    elif call.data == 'upload_markscard_pdf':
        handle_upload_markscard_pdf(call.message)
    elif call.data == 'sgpa':
        handle_sgpa(call.message)
    elif call.data == 'cgpa':
        handle_cgpa(call.message)
    elif call.data == 'profile':
        handle_profile(call.message)
    elif call.data == 'update_profile':
        handle_update_profile(call.message)
    elif call.data == 'generate_report':
        handle_generate_report(call.message)
    elif call.data == 'set_reminder':
        handle_set_reminder(call.message)
    elif call.data == 'share_document':
        handle_share_document(call.message)
    elif call.data == 'list_resources':
        handle_list_resources(call.message)
    elif call.data == 'job_opportunities':
        handle_job_opportunities(call.message)
    elif call.data == 'feedback':
        handle_feedback(call.message)
    elif call.data == 'logout':
        handle_logout(call.message)

@bot.message_handler(commands=['register'])
def handle_register(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id:
        bot.send_message(chat_id, 'Please logout first using /logout before registering a new account.')
        return
    user_sessions[chat_id]['state'] = states['USERNAME']
    bot.send_message(chat_id, 'Enter your username:')

@bot.message_handler(commands=['login'])
def handle_login(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id:
        bot.send_message(chat_id, 'Please logout first using /logout before logging in.')
        return
    user_sessions[chat_id]['state'] = states['LOGIN_USERNAME']
    bot.send_message(chat_id, 'Enter your username:')

@bot.message_handler(commands=['sgpa'])
def handle_sgpa(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT sgpa FROM users WHERE user_id = %s', (user_id,))
            sgpa = cur.fetchone()
            cur.close()
            close_db_connection(conn)

            if sgpa and sgpa[0] is not None:
                bot.send_message(chat_id, f'Your SGPA is: {sgpa[0]:.2f}')
            else:
                bot.send_message(chat_id, 'No SGPA records found. Please upload your marks card using /upload_markscard_pdf.')
        except Exception as e:
            bot.send_message(chat_id, f'Error fetching SGPA: {e}')
            close_db_connection(conn)

@bot.message_handler(commands=['cgpa'])
def handle_cgpa(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT sgpa FROM users WHERE user_id = %s', (user_id,))
            rows = cur.fetchall()
            cur.close()
            close_db_connection(conn)

            if rows:
                sgpa_values = [row[0] for row in rows]
                total_sgpa = sum(sgpa_values)
                num_semesters = len(sgpa_values)
                cgpa = total_sgpa / num_semesters if num_semesters > 0 else 0

                if num_semesters == 1:
                    cgpa = sgpa_values[0]

                # Update CGPA in users table
                conn = get_db_connection()
                if conn:
                    try:
                        cur = conn.cursor()
                        cur.execute('UPDATE users SET cgpa = %s WHERE user_id = %s', (cgpa, user_id))
                        conn.commit()
                        cur.close()
                        close_db_connection(conn)
                    except Exception as e:
                        bot.send_message(chat_id, f'Error updating CGPA: {e}')
                        close_db_connection(conn)

                bot.send_message(chat_id, f'Your CGPA is: {cgpa:.2f}')
            else:
                bot.send_message(chat_id, 'No SGPA records found. Please upload your marks card using /upload_markscard_pdf.')
        except Exception as e:
            bot.send_message(chat_id, f'Error calculating CGPA: {e}')
            close_db_connection(conn)

@bot.message_handler(commands=['profile'])
def handle_profile(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT full_name, semester, college, mobile, branch, year_scheme, sgpa, cgpa FROM users WHERE user_id = %s', (user_id,))
            user = cur.fetchone()
            cur.close()
            close_db_connection(conn)

            if user:
                full_name, semester, college, mobile, branch, year_scheme, sgpa, cgpa = user
                profile_message = f"""
                *Profile Information*
                Full Name: {full_name}
                Semester: {semester}
                College: {college}
                Mobile: {mobile}
                Branch: {branch}
                Year Scheme: {year_scheme}
                SGPA: {f'{sgpa:.2f}' if sgpa is not None else 'N/A'}
                CGPA: {f'{cgpa:.2f}' if cgpa is not None else 'N/A'}
                """
                bot.send_message(chat_id, profile_message, parse_mode='Markdown')
            else:
                bot.send_message(chat_id, 'Profile not found.')
        except Exception as e:
            bot.send_message(chat_id, f'Error fetching profile: {e}')
            close_db_connection(conn)

def handle_update_profile(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    markup = types.InlineKeyboardMarkup(row_width=2)
    markup.add(types.InlineKeyboardButton("Full Name", callback_data='update_full_name'),
               types.InlineKeyboardButton("Semester", callback_data='update_semester'),
               types.InlineKeyboardButton("College", callback_data='update_college'),
               types.InlineKeyboardButton("Mobile", callback_data='update_mobile'),
               types.InlineKeyboardButton("Branch", callback_data='update_branch'),
               types.InlineKeyboardButton("Year Scheme", callback_data='update_year_scheme'))
    bot.send_message(chat_id, 'Choose the information you want to update:', reply_markup=markup)
    user_sessions[chat_id]['state'] = states['UPDATE_PROFILE']

@bot.callback_query_handler(func=lambda call: call.data.startswith('update_'))
def handle_update_field(call):
    chat_id = call.message.chat.id
    field = call.data.split('_')[1]
    user_sessions[chat_id]['update_field'] = field
    user_sessions[chat_id]['state'] = states['UPDATE_PROFILE_FIELD']  # Correctly set the state
    bot.send_message(chat_id, f'Enter your new {field.replace("_", " ")}:')

@bot.message_handler(func=lambda message: user_sessions[message.chat.id]['state'] == states['UPDATE_PROFILE_FIELD'], content_types=['text'])
def handle_update_value(message):
    chat_id = message.chat.id
    field = user_sessions[chat_id]['update_field']
    user_id = user_sessions[chat_id]['userId']
    new_value = message.text

    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute(f'UPDATE users SET {field} = %s WHERE user_id = %s', (new_value, user_id))
            conn.commit()
            cur.close()
            close_db_connection(conn)
            bot.send_message(chat_id, f'{field.replace("_", " ").capitalize()} updated successfully!')
        except Exception as e:
            bot.send_message(chat_id, f'Error updating {field.replace("_", " ")}: {e}')
            close_db_connection(conn)
        finally:
            user_sessions[chat_id]['state'] = None
            user_sessions[chat_id].pop('update_field', None)

def fetch_uploaded_documents(user_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT file_id, file_name, mime_type FROM shared_documents WHERE user_id = %s', (user_id,))
            documents = cur.fetchall()
            cur.close()
            close_db_connection(conn)
            return documents
        except Exception as e:
            print(f"Error fetching documents: {e}")
            close_db_connection(conn)
            return []

@bot.message_handler(commands=['upload_markscard_pdf'])
def handle_upload_markscard_pdf(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return
    user_sessions[chat_id]['state'] = states['MARKSCARD_PDF']
    bot.send_message(chat_id, 'Please upload your marks card PDF.')

@bot.message_handler(func=lambda message: True, content_types=['text'])
def handle_text(message):
    chat_id = message.chat.id
    init_session(chat_id)
    state = user_sessions[chat_id]['state']

    if state == states['USERNAME']:
        user_sessions[chat_id]['username'] = message.text
        bot.send_message(chat_id, 'Enter your password:')
        user_sessions[chat_id]['state'] = states['PASSWORD']
    elif state == states['PASSWORD']:
        user_sessions[chat_id]['password'] = message.text
        bot.send_message(chat_id, 'Enter your full name:')
        user_sessions[chat_id]['state'] = states['FULL_NAME']
    elif state == states['FULL_NAME']:
        user_sessions[chat_id]['full_name'] = message.text
        username = user_sessions[chat_id]['username']
        conn = get_db_connection()
        if conn:
            try:
                cur = conn.cursor()
                cur.execute('SELECT user_id FROM users WHERE username = %s', (username,))
                existing_user = cur.fetchone()
                if existing_user:
                    bot.send_message(chat_id, 'Username already exists. Please login or choose a different username.')
                    user_sessions[chat_id]['state'] = states['USERNAME']
                else:
                    bot.send_message(chat_id, 'Enter your semester:')
                    user_sessions[chat_id]['state'] = states['SEMESTER']
            except Exception as e:
                bot.send_message(chat_id, f'Error during registration: {e}')
            finally:
                cur.close()
                close_db_connection(conn)
    elif state == states['SEMESTER']:
        user_sessions[chat_id]['semester'] = message.text
        bot.send_message(chat_id, 'Enter your college name:')
        user_sessions[chat_id]['state'] = states['COLLEGE']
    elif state == states['COLLEGE']:
        user_sessions[chat_id]['college'] = message.text
        bot.send_message(chat_id, 'Enter your mobile number:')
        user_sessions[chat_id]['state'] = states['MOBILE']
    elif state == states['MOBILE']:
        mobile_number = message.text
        if len(mobile_number) != 10 or not mobile_number.isdigit():
            bot.send_message(chat_id, 'Invalid mobile number. Please enter a 10-digit mobile number:')
        else:
            user_sessions[chat_id]['mobile'] = mobile_number
            bot.send_message(chat_id, 'Enter your branch:')
            user_sessions[chat_id]['state'] = states['BRANCH']
    elif state == states['BRANCH']:
        user_sessions[chat_id]['branch'] = message.text
        bot.send_message(chat_id, 'Enter your year scheme:')
        user_sessions[chat_id]['state'] = states['YEAR_SCHEME']
    elif state == states['YEAR_SCHEME']:
        year_scheme = message.text
        full_name = user_sessions[chat_id]['full_name']
        username = user_sessions[chat_id]['username']
        password = hash_password(user_sessions[chat_id]['password'])
        semester = user_sessions[chat_id]['semester']
        college = user_sessions[chat_id]['college']
        mobile = user_sessions[chat_id]['mobile']
        branch = user_sessions[chat_id]['branch']
        
        conn = get_db_connection()
        if conn:
            try:
                cur = conn.cursor()
                cur.execute('INSERT INTO users (full_name, username, password, semester, college, mobile, branch, year_scheme, sgpa, cgpa, chat_id) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING user_id', 
                            (full_name, username, password, semester, college, mobile, branch, year_scheme, None, None, chat_id))
                user_id = cur.fetchone()[0]
                conn.commit()
                user_sessions[chat_id]['userId'] = user_id
                bot.send_message(chat_id, 'Registration successful! You can now use the menu to navigate.')
                user_sessions[chat_id]['state'] = None
            except Exception as e:
                bot.send_message(chat_id, f'Error during registration: {e}')
            finally:
                cur.close()
                close_db_connection(conn)
    elif state == states['LOGIN_USERNAME']:
        user_sessions[chat_id]['username'] = message.text
        bot.send_message(chat_id, 'Enter your password:')
        user_sessions[chat_id]['state'] = states['LOGIN_PASSWORD']
    elif state == states['LOGIN_PASSWORD']:
        provided_password = message.text
        username = user_sessions[chat_id]['username']
        
        conn = get_db_connection()
        if conn:
            try:
                cur = conn.cursor()
                cur.execute('SELECT user_id, password FROM users WHERE username = %s', (username,))
                user = cur.fetchone()
                if user and check_password(user[1].tobytes(), provided_password):  # Convert stored password to bytes
                    user_sessions[chat_id]['userId'] = user[0]
                    bot.send_message(chat_id, 'Login successful! You can now use the menu to navigate.')
                    user_sessions[chat_id]['state'] = None
                else:
                    bot.send_message(chat_id, 'Invalid username or password. Please try again.')
                    user_sessions[chat_id]['state'] = states['LOGIN_USERNAME']
            except Exception as e:
                bot.send_message(chat_id, f'Error during login: {e}')
            finally:
                cur.close()
                close_db_connection(conn)
    elif state == states['RESET_PASSWORD']:
        new_password = hash_password(message.text)
        username = user_sessions[chat_id]['username']
        
        conn = get_db_connection()
        if conn:
            try:
                cur = conn.cursor()
                cur.execute('UPDATE users SET password = %s WHERE username = %s', (new_password, username))
                conn.commit()
                bot.send_message(chat_id, 'Password reset successfully!')
                user_sessions[chat_id]['state'] = None
            except Exception as e:
                bot.send_message(chat_id, f'Error resetting password: {e}')
            finally:
                cur.close()
                close_db_connection(conn)
    elif state == states['REMINDER_TIME']:
        user_sessions[chat_id]['reminder_time'] = message.text
        bot.send_message(chat_id, 'Enter the reminder message:')
        user_sessions[chat_id]['state'] = states['REMINDER_MESSAGE']
    elif state == states['REMINDER_MESSAGE']:
        reminder_message = message.text
        reminder_time = user_sessions[chat_id]['reminder_time']
        user_id = user_sessions[chat_id]['userId']

        if add_reminder(user_id, reminder_time, reminder_message):
            bot.send_message(chat_id, 'Reminder set successfully!')
        else:
            bot.send_message(chat_id, 'Error setting reminder.')

        user_sessions[chat_id]['state'] = None
    elif state == states['FEEDBACK']:
        feedback_text = message.text
        user_id = user_sessions[chat_id]['userId']
        if save_feedback(user_id, feedback_text):
            bot.send_message(chat_id, 'Thank you for your feedback!')
        else:
            bot.send_message(chat_id, 'Error saving feedback.')
        user_sessions[chat_id]['state'] = None
    else:
        bot.send_message(chat_id, 'Unknown command. Please use /menu to see available options.')

@bot.message_handler(content_types=['document', 'photo'])
def handle_document(message):
    chat_id = message.chat.id
    init_session(chat_id)
    state = user_sessions[chat_id]['state']

    if user_sessions[chat_id]['userId'] is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    if state == states['MARKSCARD_PDF']:
        if message.content_type == 'document' and message.document.mime_type == 'application/pdf':
            file_id = message.document.file_id
            user_id = user_sessions[chat_id]['userId']
            
            # Check if the file already exists
            if check_existing_marks_card(user_id, file_id):
                sgpa = fetch_sgpa(user_id)
                bot.send_message(chat_id, f'You have already uploaded this marks card. Your SGPA is: {sgpa:.2f}')
                return
            
            file_info = bot.get_file(file_id)
            file_url = f'https://api.telegram.org/file/bot{BOT_TOKEN}/{file_info.file_path}'

            response = requests.get(file_url)

            with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as pdf_file:
                pdf_file.write(response.content)
                pdf_path = pdf_file.name

            excel_path = pdf_path.replace('.pdf', '.xlsx')

            # Convert PDF to Excel
            document = ap.Document(pdf_path)
            save_option = ap.ExcelSaveOptions()
            document.save(excel_path, save_option)

            # Process Excel data
            sgpa = process_excel_data(excel_path)
            
            # Save SGPA to database
            save_sgpa_to_db(user_id, sgpa)
            
            # Save the marks card to the database
            save_marks_card(user_id, file_id)
            
            bot.send_message(chat_id, 'Marks card PDF uploaded and processed successfully. SGPA has been updated.')
            user_sessions[chat_id]['state'] = None
        else:
            bot.send_message(chat_id, 'Unsupported file format. Please upload a PDF file.')
    elif state == states['SHARE_DOCUMENT']:
        if message.content_type in ['document', 'photo']:
            file_id = message.document.file_id if message.content_type == 'document' else message.photo[-1].file_id
            file_name = message.document.file_name if message.content_type == 'document' else 'photo.jpg'
            mime_type = message.document.mime_type if message.content_type == 'document' else 'image/jpeg'
            user_id = user_sessions[chat_id]['userId']

            if save_shared_document(user_id, file_id, file_name, mime_type):
                bot.send_message(chat_id, f'Document {file_name} shared successfully!')
            else:
                bot.send_message(chat_id, 'Error sharing document.')

def check_existing_marks_card(user_id, file_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT card_id FROM marks_cards WHERE user_id = %s AND file_id = %s', (user_id, file_id))
            existing_card = cur.fetchone()
            cur.close()
            close_db_connection(conn)
            return existing_card is not None
        except Exception as e:
            print(f"Error checking existing marks card: {e}")
            close_db_connection(conn)
            return False

def fetch_sgpa(user_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT sgpa FROM users WHERE user_id = %s', (user_id,))
            sgpa = cur.fetchone()[0]
            cur.close()
            close_db_connection(conn)
            return sgpa
        except Exception as e:
            print(f"Error fetching SGPA: {e}")
            close_db_connection(conn)
            return None

def save_marks_card(user_id, file_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('INSERT INTO marks_cards (user_id, file_id) VALUES (%s, %s)', (user_id, file_id))
            conn.commit()
            cur.close()
            close_db_connection(conn)
            return True
        except Exception as e:
            print(f"Error saving marks card: {e}")
            close_db_connection(conn)
            return False

def process_excel_data(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    total_points = 0
    total_credits = 0

    # Adjust the expected number of columns based on the Excel file's structure
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if len(row) >= 4:  # Ensure there are at least 4 values in the row
            subject_code, subject_name, internal_marks, external_marks = row[:4]
            
            # Ensure internal_marks and external_marks are not None and convert them to integers
            if internal_marks is None or not isinstance(internal_marks, (int, float)):
                internal_marks = 0
            if external_marks is None or not isinstance(external_marks, (int, float)):
                external_marks = 0
            
            total_marks = int(internal_marks) + int(external_marks)
            grade_points = convert_to_grade_points(total_marks)
            credits = get_credits_for_subject(subject_code)
            total_points += grade_points * credits
            total_credits += credits

    sgpa = total_points / total_credits if total_credits != 0 else 0
    return sgpa

def convert_to_grade_points(total_marks):
    if total_marks >= 90:
        return 10
    elif total_marks >= 80:
        return 9
    elif total_marks >= 70:
        return 8
    elif total_marks >= 60:
        return 7
    elif total_marks >= 50:
        return 6
    elif total_marks >= 40:
        return 5
    else:
        return 0

def get_credits_for_subject(subject_code):
    credits_map = {
        #5th sem 21 batch
       '21CS51': 3,'21CSL582': 1,'21CS52': 4,'21CS53': 3,'21CS54': 3,'21CSL55': 1,'21RMI56': 2,'21CIV57': 1,
       #3rd sem 21 batch
        '21MAT31':3,'21CS382':1,'21CS32':4,'21CS33':4,'21CS34':3,'21CSL35':1,'21SCR36':1,'21KBK37':1,
        #3rd sem 22 batch
        'BCS301':4,'BCS302':4,'BCS303':4,'BCS304':3,'BCSL305':1,'BSCK307':1,'BNSK359':0,'BCS306A':3,'BCS358C':1,
        #1st sem 22 batch
        'BMATS101':4,'BPHYS102':4,'BPOPS103':3,'BESCK104B':3,'BETCK105I':3,'BENGK106':1,'BICOK107':1,'BIDTK158':1,
    }
    return credits_map.get(subject_code, 0)

def save_sgpa_to_db(user_id, sgpa):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('UPDATE users SET sgpa = %s WHERE user_id = %s', (sgpa, user_id))
            conn.commit()
            cur.close()
            close_db_connection(conn)
        except Exception as e:
            print(f"Error saving SGPA to database: {e}")
            close_db_connection(conn)
            
def save_marks_to_db(user_id, subject_code, subject_name, internal_marks, external_marks, sgpa, credits):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('INSERT INTO marks (user_id, subject_code, subject_name, internal_marks, external_marks, total, sgpa, credits) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)',
                        (user_id, subject_code, subject_name, internal_marks, external_marks, internal_marks + external_marks, sgpa, credits))
            conn.commit()
            cur.close()
            close_db_connection(conn)
            return True
        except Exception as e:
            print(f"Error saving marks to database: {e}")
            close_db_connection(conn)
            return False
    return False
@bot.message_handler(commands=['reset_password'])
def handle_reset_password(message):
    chat_id = message.chat.id
    init_session(chat_id)
    bot.send_message(chat_id, 'Enter your username:')
    user_sessions[chat_id]['state'] = states['LOGIN_USERNAME']

@bot.message_handler(func=lambda message: user_sessions[message.chat.id]['state'] == states['LOGIN_USERNAME'], content_types=['text'])
def handle_username_for_reset(message):
    chat_id = message.chat.id
    username = message.text
    user_sessions[chat_id]['username'] = username
    bot.send_message(chat_id, 'Enter your new password:')
    user_sessions[chat_id]['state'] = states['RESET_PASSWORD']

@bot.message_handler(func=lambda message: user_sessions[message.chat.id]['state'] == states['RESET_PASSWORD'], content_types=['text'])
def handle_new_password(message):
    chat_id = message.chat.id
    new_password = hash_password(message.text)
    username = user_sessions[chat_id]['username']
    
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('UPDATE users SET password = %s WHERE username = %s', (new_password, username))
            conn.commit()
            bot.send_message(chat_id, 'Password reset successfully!')
            user_sessions[chat_id]['state'] = None
        except Exception as e:
            bot.send_message(chat_id, f'Error resetting password: {e}')
        finally:
            cur.close()
            close_db_connection(conn)

def handle_logout(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_sessions[chat_id] = {'state': None, 'username': None, 'userId': None}
    bot.send_message(chat_id, 'You have been logged out successfully.')

def generate_report(user_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT full_name, semester, college, branch, sgpa, cgpa FROM users WHERE user_id = %s', (user_id,))
            user = cur.fetchone()
            cur.close()
            close_db_connection(conn)

            report_path = f'report_{user_id}.pdf'
            c = SimpleDocTemplate(report_path, pagesize=letter)

            styles = getSampleStyleSheet()
            title_style = styles['Title']
            title = Paragraph('Campus Connect', title_style)

            table_data = [
                ['Field', 'Details'],
                ['Full Name', user[0]],
                ['Semester', user[1]],
                ['College', user[2]],
                ['Branch', user[3]],
                ['SGPA', f'{user[4]:.2f}'],
                ['CGPA', f'{user[5]:.2f}'],
            ]

            table = Table(table_data, colWidths=[150, 350])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 14),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ]))

            elements = [title, table]
            c.build(elements)
            return report_path
        except Exception as e:
            logging.error(f"Error generating report: {e}")
            close_db_connection(conn)
            return None
    else:
        return None
@bot.message_handler(commands=['generate_report'])
def handle_generate_report(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    report_path = generate_report(user_id)
    if report_path:
        with open(report_path, 'rb') as report_file:
            bot.send_document(chat_id, report_file)
    else:
        bot.send_message(chat_id, 'Error generating report.')

def add_reminder(user_id, time_str, message):
    job_id = str(uuid4())
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('INSERT INTO reminders (user_id, time_str, message, job_id) VALUES (%s, %s, %s, %s)', 
                        (user_id, time_str, message, job_id))
            conn.commit()
            cur.close()
            close_db_connection(conn)

            hour, minute = map(int, time_str.split(':'))
            scheduler.add_job(send_reminder, CronTrigger(hour=hour, minute=minute), args=[user_id, message], id=job_id)
            return True
        except Exception as e:
            print(f"Error adding reminder: {e}")
            close_db_connection(conn)
            return False
    return False

def send_reminder(user_id, message):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT chat_id FROM users WHERE user_id = %s', (user_id,))
            chat_id = cur.fetchone()[0]
            cur.close()
            close_db_connection(conn)

            bot.send_message(chat_id, f"Reminder: {message}")
        except Exception as e:
            print(f"Error sending reminder: {e}")
            close_db_connection(conn)

def schedule_reminders():
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT job_id, user_id, time_str, message FROM reminders')
            reminders = cur.fetchall()
            cur.close()
            close_db_connection(conn)

            for job_id, user_id, time_str, message in reminders:
                hour, minute = map(int, time_str.split(':'))
                scheduler.add_job(send_reminder, CronTrigger(hour=hour, minute=minute), args=[user_id, message], id=job_id)
        except Exception as e:
            print(f"Error scheduling reminders: {e}")
            close_db_connection(conn)

@bot.message_handler(commands=['set_reminder'])
def handle_set_reminder(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    msg = bot.send_message(chat_id, 'Enter the reminder time in HH:MM format:')
    bot.register_next_step_handler(msg, get_reminder_time)

def get_reminder_time(message):
    chat_id = message.chat.id
    reminder_time = message.text
    msg = bot.send_message(chat_id, 'Enter the reminder message:')
    bot.register_next_step_handler(msg, get_reminder_message, reminder_time)

def get_reminder_message(message, reminder_time):
    chat_id = message.chat.id
    reminder_message = message.text
    user_id = user_sessions[chat_id]['userId']

    if add_reminder(user_id, reminder_time, reminder_message):
        bot.send_message(chat_id, 'Reminder set successfully!')
        schedule_reminders()
    else:
        bot.send_message(chat_id, 'Error setting reminder.')

scheduler = BackgroundScheduler()
scheduler.start()

schedule_reminders()

@bot.message_handler(commands=['job_opportunities'])
def handle_job_opportunities(message):
    chat_id = message.chat.id
    init_session(chat_id)

    job_opportunities = fetch_job_opportunities()
    if job_opportunities:
        for job in job_opportunities:
            job_message = f"**{job[0]}** at **{job[1]}**\n{job[2]}\n[More Info]({job[3]})"
            bot.send_message(chat_id, job_message, parse_mode='Markdown', disable_web_page_preview=True)
    else:
        bot.send_message(chat_id, 'No job opportunities available.')

def fetch_job_opportunities():
    # Replace with actual API or web scraping logic to fetch job opportunities
    # Example job opportunity data
    job_opportunities = [
        ("Software Engineer Intern", "Google", "An exciting opportunity to work on cutting-edge technology.", "https://careers.google.com/jobs/results/"),
        ("Data Analyst", "Facebook", "Analyze data to drive key business decisions.", "https://www.facebook.com/careers/jobs"),
        ("Backend Developer", "Amazon", "Develop and maintain backend services.", "https://www.amazon.jobs/en/"),
    ]
    return job_opportunities

@bot.message_handler(commands=['share_document'])
def handle_share_document(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    bot.send_message(chat_id, 'Upload the document you want to share:')
    user_sessions[chat_id]['state'] = states['SHARE_DOCUMENT']

def save_shared_document(user_id, file_id, file_name, mime_type):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('INSERT INTO shared_documents (user_id, file_id, file_name, mime_type) VALUES (%s, %s, %s, %s)', 
                        (user_id, file_id, file_name, mime_type))
            conn.commit()
            cur.close()
            close_db_connection(conn)
            return True
        except Exception as e:
            print(f"Error saving shared document: {e}")
            close_db_connection(conn)
            return False
    return False

@bot.message_handler(commands=['list_resources'])
def handle_list_resources(message):
    chat_id = message.chat.id
    init_session(chat_id)
    user_id = user_sessions[chat_id]['userId']
    if user_id is None:
        bot.send_message(chat_id, 'Please login first using /login.')
        return

    resources = fetch_resources(user_id)
    if resources:
        for resource in resources:
            file_id, file_name, mime_type = resource
            if mime_type == 'application/pdf':
                bot.send_document(chat_id, file_id, caption=file_name)
            elif mime_type.startswith('image/'):
                bot.send_photo(chat_id, file_id, caption=file_name)
            else:
                bot.send_message(chat_id, f"{file_name} - Shared document")
    else:
        bot.send_message(chat_id, 'No resources available.')

def fetch_resources(user_id):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('SELECT file_id, file_name, mime_type FROM shared_documents WHERE user_id = %s', (user_id,))
            resources = cur.fetchall()
            cur.close()
            close_db_connection(conn)
            return resources
        except Exception as e:
            print(f"Error fetching resources: {e}")
            close_db_connection(conn)
            return []

@bot.message_handler(commands=['feedback'])
def handle_feedback(message):
    chat_id = message.chat.id
    init_session(chat_id)

    bot.send_message(chat_id, 'Enter your feedback:')
    user_sessions[chat_id]['state'] = states['FEEDBACK']

@bot.message_handler(func=lambda message: user_sessions[message.chat.id]['state'] == states['FEEDBACK'], content_types=['text'])
def handle_feedback_message(message):
    chat_id = message.chat.id
    user_id = user_sessions[chat_id]['userId']
    feedback = message.text

    if save_feedback(user_id, feedback):
        bot.send_message(chat_id, 'Thank you for your feedback!')
    else:
        bot.send_message(chat_id, 'Error saving feedback.')

    user_sessions[chat_id]['state'] = None

def save_feedback(user_id, feedback):
    conn = get_db_connection()
    if conn:
        try:
            cur = conn.cursor()
            cur.execute('INSERT INTO feedback (user_id, feedback_text) VALUES (%s, %s)', (user_id, feedback))
            conn.commit()
            cur.close()
            close_db_connection(conn)
            return True
        except Exception as e:
            print(f"Error saving feedback: {e}")
            close_db_connection(conn)
            return False
    return False

def start_polling():
    while True:
        try:
            bot.polling(none_stop=True, interval=0, timeout=60)
        except Exception as e:
            print(f"Error occurred: {e}")
            time.sleep(15)

if __name__ == "__main__":
    start_polling()
